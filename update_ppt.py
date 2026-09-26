from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
import os

try:
    ppt_path = "C:\\Users\\P52\\Desktop\\vuln-shield\\VulnShield_Spectra26.pptx"
    output_path = "C:\\Users\\P52\\Desktop\\vuln-shield\\VulnShield_Presentation.pptx"
    
    if not os.path.exists(ppt_path):
        ppt_path = "VulnShield_Spectra26.pptx"
        output_path = "VulnShield_Presentation.pptx"

    prs = Presentation(ppt_path)
    
    # 1 is the layout index for Title and Content in most themes
    slide_layout = prs.slide_layouts[1] 
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    title.text = "CONCLUSION: STOPPING THE NEXT CYBER ATTACK"
    
    for shape in slide.placeholders:
        if shape.is_placeholder and shape.placeholder_format.idx == 1:
            body = shape
            break
            
    tf = body.text_frame
    tf.text = "📉 The Threat: Hackers no longer attack servers; they attack our personal mobile phones."
    
    p = tf.add_paragraph()
    p.text = "🛡️ The Defense: VulnShield acts as an instant, zero-agent digital bodyguard to detect spyware and web vulnerabilities in real-time."
    
    p = tf.add_paragraph()
    p.text = "🚀 The Impact: Secures digital identities, prevents data leaks, and empowers users to fight back against modern cyber threats."

    # Try to make it look decent
    for p in tf.paragraphs:
        p.font.size = Pt(20)

    prs.save(output_path)
    print(f"SUCCESS: Saved new presentation to {output_path}")

except Exception as e:
    print(f"ERROR: {str(e)}")
